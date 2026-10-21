from PyQt5 import QtCore, QtWidgets, QtGui
import openpyxl


from db_API import Database as DB
from Data_window import DataWindow
from Doc_Window import DocWindow
from Scan_window import ScanWindow
from InsertWindow import InsertWindow, ChangeWindow
from Fizo_window import FizoWindow
import Tab_widgets2
import datetime
import style


class _WideTabBar(QtWidgets.QTabBar):
    """Панель вкладок, которая никогда не даёт Qt обрезать заголовки:
    переопределяем tabSizeHint — ширина = реальный текст жирным шрифтом
    (как в выбранной вкладке) + запас на padding из QSS (2*14px) и рамки."""

    EXTRA_WIDTH = 60

    def tabSizeHint(self, index):
        size = super().tabSizeHint(index)
        base = QtGui.QFont(self.font())
        base.setPointSize(12)              # font-size: 12pt из QSS
        bold = QtGui.QFont(base)
        bold.setBold(True)                 # выбранная вкладка — жирная
        needed = QtGui.QFontMetrics(bold).horizontalAdvance(self.tabText(index)) \
                 + self.EXTRA_WIDTH
        if size.width() < needed:
            size.setWidth(needed)
        return size


class MainTab(QtWidgets.QTabWidget):
    def __init__(self, data, parent=None):
        super().__init__(parent=parent)

        self.db = parent.db

        # собственная панель ярлыков: заголовки («Документы», «Члены семьи»)
        # всегда помещаются целиком и не обрезаются до «кумен»
        self.setTabBar(_WideTabBar(self))
        self.tabBar().setExpanding(False)
        self.tabBar().setElideMode(QtCore.Qt.ElideNone)
        self.tabBar().setUsesScrollButtons(False)

        if data:
            data = self.db.get_data(data)[0]
            person_id = data[0]

            person_w = Tab_widgets2.PersonWidget(self, data)
            document_w = Tab_widgets2.DocsWidget(self, person_id)
            family_w = Tab_widgets2.FamilyWidget(self, person_id)
            service_w = Tab_widgets2.ServiceWidget(self, person_id)
            fizo_w = Tab_widgets2.FIZOWidget(self, person_id)


            self.addTab(person_w, 'Главная')
            self.addTab(document_w, 'Документы')
            self.addTab(family_w, 'Члены семьи')
            self.addTab(service_w, 'Прохождение службы')

            if data[12] == 'True':
                self.addTab(fizo_w, 'ФИЗО')

        else:
            main_layout = QtWidgets.QVBoxLayout()
            main_layout.setSpacing(18)
            logo_w = QtWidgets.QLabel('')
            pixmap = QtGui.QPixmap('./data/logo.png')
            logo_w.setPixmap(pixmap.scaled(160, 160,
                                           QtCore.Qt.KeepAspectRatio,
                                           QtCore.Qt.SmoothTransformation))
            logo_w.setAlignment(QtCore.Qt.AlignHCenter | QtCore.Qt.AlignCenter)
            main_layout.addWidget(logo_w)
            logo_label = QtWidgets.QLabel(
                '<div align="center">'
                '<span style="font-size:30pt;font-weight:700;color:{};">'
                'СПО учета личного состава</span><br>'
                '<span style="font-size:16pt;color:{};">'
                'Выберите сотрудника из списка слева или добавьте новый профиль'
                '</span></div>'.format(style.ACCENT, style.MUTED))
            logo_label.setStyleSheet('background: transparent; border: none;')
            main_layout.addWidget(logo_label)
            main_layout.setAlignment(QtCore.Qt.AlignCenter | QtCore.Qt.AlignHCenter)
            self.setLayout(main_layout)


class MainWindow(QtWidgets.QMainWindow):
    def __init__(self, db, parent=None):
        super().__init__(parent=parent)
        
        self.db = DB(db)

        # при каждом запуске приложения — светлая тема (новый интерфейс)
        self._dark_on = False
        style.set_dark_mode(False)
        style.apply_theme(QtWidgets.QApplication.instance())

        self.tool_bar = QtWidgets.QToolBar()
        # тулбар — горизонтальный, сверху окна (привычнее и компактнее по высоте)
        self.tool_bar.setOrientation(QtCore.Qt.Horizontal)
        self.tool_bar.setIconSize(QtCore.QSize(32, 32))
        self.tool_bar.setToolButtonStyle(QtCore.Qt.ToolButtonTextUnderIcon)
        self.tool_bar.setMovable(False)
        self.tool_bar.setFloatable(False)

        insert_profile = QtWidgets.QAction(QtGui.QIcon('./data/add.png'), 'Добавить', self)
        insert_profile.triggered.connect(self.insert)

        delete_profile = QtWidgets.QAction(QtGui.QIcon('./data/delete.png'), 'Удалить', self)
        delete_profile.triggered.connect(self.delete_profile)

        update_profile = QtWidgets.QAction(QtGui.QIcon('./data/edit.png'), 'Редактировать', self)
        update_profile.triggered.connect(self.change_data)

        add_scan = QtWidgets.QAction(QtGui.QIcon('./data/folder.png'), 'Прикрепить\nдокумент', self)
        add_scan.triggered.connect(self.add_scan)
        
        data_menu = QtWidgets.QAction(QtGui.QIcon('./data/addressbook.png'), 'Cловарь', self)
        data_menu.triggered.connect(self.data_change)

        fizo_menu = QtWidgets.QAction(QtGui.QIcon('./data/sport.png'), 'ФИЗО', self)
        fizo_menu.triggered.connect(self.fizo_open)
        
        im_export_menu = QtWidgets.QAction(QtGui.QIcon('./data/export.png'), 'Экспорт', self)
        im_export_menu.triggered.connect(self.export)

        import_menu = QtWidgets.QAction(QtGui.QIcon('./data/folder.png'), 'Импорт\nданных', self)
        import_menu.setToolTip('Перенести данные из базы (.db) старой версии. '
                               'Текущая база будет очищена (перед этим создан бекап).')
        import_menu.triggered.connect(self.import_db)

        docs_menu = QtWidgets.QAction(QtGui.QIcon('./data/data.png'), 'Формы\nдокументов', self)
        docs_menu.triggered.connect(self.create_documents)

        # --- кнопка переключения чёрной темы (добавляется в конец тулбара) ---

        def _theme_icon(dark):
            """Рисунок «полусфера день/ночь» — не зависит от внешних файлов."""
            pm = QtGui.QPixmap(24, 24)
            pm.fill(QtCore.Qt.transparent)
            p = QtGui.QPainter(pm)
            p.setRenderHint(QtGui.QPainter.Antialiasing)
            pen = QtGui.QPen(QtGui.QColor('#f0b429' if dark else '#5b6b7f'))
            pen.setWidth(2)
            p.setPen(pen)
            p.setBrush(QtCore.Qt.NoBrush)
            p.drawEllipse(3, 3, 18, 18)
            path = QtGui.QPainterPath()
            path.moveTo(12, 3)
            path.arcTo(QtCore.QRectF(3, 3, 18, 18), 90, -180)
            path.closeSubpath()
            p.setBrush(QtGui.QColor('#f0b429' if dark else '#5b6b7f'))
            p.drawPath(path)
            p.end()
            return QtGui.QIcon(pm)

        self._dark_on = False
        self.theme_action = QtWidgets.QAction(_theme_icon(False), 'Чёрная тема', self)
        self.theme_action.setToolTip('Включить/выключить чёрную тему СПО')
        self.theme_action.triggered.connect(self.toggle_dark_theme)

        self.tool_bar.addAction(insert_profile)
        self.tool_bar.addAction(update_profile)
        self.tool_bar.addSeparator()
        self.tool_bar.addAction(add_scan)
        self.tool_bar.addAction(delete_profile)
        self.tool_bar.addSeparator()
        self.tool_bar.addAction(data_menu)
        self.tool_bar.addSeparator()
        self.tool_bar.addAction(docs_menu)
        self.tool_bar.addSeparator()
        self.tool_bar.addAction(fizo_menu)
        self.tool_bar.addSeparator()
        self.tool_bar.addAction(im_export_menu)
        self.tool_bar.addAction(import_menu)
        self.tool_bar.addSeparator()

        # кнопка темы — в конце горизонтального тулбара, текст сбоку от иконки
        self.tool_bar.addAction(self.theme_action)
        theme_btn = self.tool_bar.widgetForAction(self.theme_action)
        if theme_btn is not None:
            theme_btn.setObjectName('theme_btn')
            theme_btn.setToolButtonStyle(QtCore.Qt.ToolButtonTextBesideIcon)

        self.tool_bar.setMovable(False)
        self.tool_bar.setToolButtonStyle(QtCore.Qt.ToolButtonTextUnderIcon)
        self.tool_bar.setIconSize(QtCore.QSize(42, 42))

        self.central_widget = QtWidgets.QWidget()
        self.central_widget.setObjectName('CentralHost')
        self.splitter = QtWidgets.QSplitter(QtCore.Qt.Horizontal)
        self.hbox = QtWidgets.QHBoxLayout()
        self.person_widget = MainTab(None, parent=self)

        self.list = QtWidgets.QListWidget()
        self.list_items = dict()
        persons = self.db.select("""
            SELECT 
            persons."ID",
            persons.second_name, 
            persons.first_name, 
            persons.patr
            FROM 
            persons
            ORDER BY
            persons.second_name ASC;
            """)
        if persons:
            for line in persons:
                if line[1] and line[2] and line[3]:
                    text = '{} {}.{}.'.format(line[1], line[2][0], line[3][0])
                self.list_items[text] = line[0]
                self.list.addItem(QtWidgets.QListWidgetItem(text))

        self.list.currentRowChanged.connect(self.row_changed)
        
        self.search_area = QtWidgets.QVBoxLayout()
        self.search_area.setSpacing(12)
        self.search_area.setContentsMargins(10, 10, 10, 10)

        search_header = QtWidgets.QLabel('Поиск сотрудников')
        search_header.setStyleSheet(
            'font-size: 16pt; font-weight: 700; color: {};'.format(style.ACCENT))
        self.search_area.addWidget(search_header)

        search_spec = QtWidgets.QGroupBox('Параметры поиска')

        self.current_check = QtWidgets.QCheckBox('Действующий')
        self.archive_check = QtWidgets.QCheckBox('Резерв')
        self.current_check.setChecked(True)
        self.current_check.stateChanged.connect(self.search_spec_changed)
        self.archive_check.stateChanged.connect(self.search_spec_changed)

        # переключатель «Военнослужащий / ГП»: два взаимоисключающих
        # radiobutton в одной группе — по умолчанию выбран «ВС».
        self.mil_radio = QtWidgets.QRadioButton('Военнослужащий')
        self.gp_radio = QtWidgets.QRadioButton('ГП ВС РФ')
        self.mil_radio.setChecked(True)
        self.mil_radio.setCursor(QtCore.Qt.PointingHandCursor)
        self.gp_radio.setCursor(QtCore.Qt.PointingHandCursor)
        self.mil_radio.setStyleSheet('font-size: 11pt;')
        self.gp_radio.setStyleSheet('font-size: 11pt;')
        self.mil_radio.toggled.connect(self.search_spec_changed)
        self.gp_radio.toggled.connect(self.search_spec_changed)

        chose_box = QtWidgets.QHBoxLayout()
        chose_box.setContentsMargins(4, 2, 4, 2)
        chose_box.setSpacing(12)
        chose_box.addWidget(self.mil_radio)
        chose_box.addWidget(self.gp_radio)
        chose_box.addStretch(1)

        h_box = QtWidgets.QHBoxLayout()
        h_box.addWidget(self.current_check)
        h_box.addWidget(self.archive_check)

        v_box = QtWidgets.QVBoxLayout()
        v_box.setSpacing(6)
        v_box.addLayout(chose_box)
        v_box.addLayout(h_box)

        search_spec.setLayout(v_box)
        
        self.search_row = QtWidgets.QHBoxLayout()
        self.search_row.setSpacing(8)
        self.search_field = QtWidgets.QLineEdit()
        self.search_field.setPlaceholderText('Фамилия…')
        self.search_field.textChanged.connect(self.find)
        self.search_btn = QtWidgets.QPushButton(QtGui.QIcon('./data/find.png'), '')
        self.search_btn.setObjectName('GhostBtn')
        self.search_btn.setFixedWidth(46)
        self.search_btn.clicked.connect(self.find)
        self.search_row.addWidget(self.search_field)
        self.search_row.addWidget(self.search_btn)
        
        self.search_update_btn = QtWidgets.QPushButton(
            QtGui.QIcon('./data/refresh.png'), ' Обновить список')
        self.search_update_btn.setObjectName('GhostBtn')
        self.search_update_btn.clicked.connect(self.find)
        
        self.search_area.addWidget(search_spec)
        self.search_area.addLayout(self.search_row)
        self.search_area.addWidget(self.list)
        self.search_area.addWidget(self.search_update_btn)
        self.search_w = QtWidgets.QWidget()
        self.search_w.setLayout(self.search_area)
        self.search_w.setMinimumWidth(260)
        self.search_w.setMaximumWidth(340)

        self.splitter.addWidget(self.search_w)
        self.splitter.addWidget(self.person_widget)
        #self.splitter.addWidget(self.scroll_area)
        self.hbox.setContentsMargins(12, 12, 12, 12)
        self.hbox.addWidget(self.splitter)
        self.central_widget.setLayout(self.hbox)

        self.splitter.setHandleWidth(12)
        self.splitter.setCollapsible(0, False)
        self.splitter.setStretchFactor(0, 0)
        self.splitter.setStretchFactor(1, 1)

        self.statusbar = QtWidgets.QStatusBar()

        self.setCentralWidget(self.central_widget)
        # тулбар сверху окна (привычнее, чем боковой)
        self.addToolBar(QtCore.Qt.TopToolBarArea, self.tool_bar)
        self.setStatusBar(self.statusbar)
        self.statusbar.showMessage('Готово')
        
        self.search_spec_query = ''
        # по умолчанию выбран «Военнослужащий» — применяем фильтр
        self.mil_radio.setChecked(True)
        self.search_spec_changed()
        
    def toggle_dark_theme(self):
        """Включает/выключает чёрную тему СПО (кнопка на тулбаре)."""
        self._dark_on = not self._dark_on
        style.set_dark_mode(self._dark_on)
        app = QtWidgets.QApplication.instance()
        style.apply_theme(app)

        # перекрасить статичные элементы, чей цвет был зашит при создании
        self.search_area.itemAt(0).widget().setStyleSheet(
            'font-size: 16pt; font-weight: 700; color: {};'.format(style.ACCENT))
        # карточка сотрудника/приветственный экран пересоздаются — подхватят новую палитру
        if self.list.currentItem():
            self.chose_person(self.list_items[self.list.currentItem().text()])
        elif self.splitter.widget(1) is not None:
            self.splitter.replaceWidget(1, MainTab(None, parent=self))

        # иконка и подпись кнопки отражают текущее состояние
        pm = QtGui.QPixmap(24, 24)
        pm.fill(QtCore.Qt.transparent)
        p = QtGui.QPainter(pm)
        p.setRenderHint(QtGui.QPainter.Antialiasing)
        col = '#f0b429' if self._dark_on else '#5b6b7f'
        pen = QtGui.QPen(QtGui.QColor(col))
        pen.setWidth(2)
        p.setPen(pen)
        p.setBrush(QtCore.Qt.NoBrush)
        p.drawEllipse(3, 3, 18, 18)
        path = QtGui.QPainterPath()
        path.moveTo(12, 3)
        path.arcTo(QtCore.QRectF(3, 3, 18, 18), 90, -180)
        path.closeSubpath()
        p.setBrush(QtGui.QColor(col))
        p.drawPath(path)
        p.end()
        self.theme_action.setIcon(QtGui.QIcon(pm))
        self.theme_action.setText('Светлая тема' if self._dark_on else 'Чёрная тема')
        self.statusbar.showMessage(
            'Чёрная тема включена' if self._dark_on else 'Светлая тема включена', 3000)

    def search_spec_changed(self):
        self.search_spec_query = ''
        if self.mil_radio.isChecked():
            self.search_spec_query += "AND persons.mil = 'True' "
        elif self.gp_radio.isChecked():
            self.search_spec_query += "AND persons.mil = 'False' "
        else:
            # ни одна кнопка не выбрана (можно снять выбор, если понадобится)
            # — не фильтруем по типу
            pass
        if self.current_check.isChecked() and not self.archive_check.isChecked():
            self.search_spec_query += "AND persons.active = 'True' "
        elif not self.current_check.isChecked() and self.archive_check.isChecked():
            self.search_spec_query += "AND persons.active = 'False' "
        elif not self.current_check.isChecked() and not self.archive_check.isChecked():
            self.search_spec_query += "AND persons.active = 'Null' "
        self.find()

        
    def row_changed(self):
        if self.list.currentItem():
            self.chose_person(self.list_items[self.list.currentItem().text()])
        
    def find(self):
        self.list.clear()
        self.list_items.clear()
        
        search = self.search_field.text().capitalize()
        
        query = """
        SELECT 
        persons."ID",
        persons.second_name, 
        persons.first_name, 
        persons.patr
        FROM 
        persons
        WHERE persons.second_name LIKE '{}%' {}
        ORDER BY
        persons.second_name ASC;        
        """.format(search, self.search_spec_query)  
        persons = self.db.select(query)     
        
        if persons:
            for line in persons:
                text = '{} {}.{}.'.format(line[1], line[2][0], line[3][0])
                self.list_items[text] = line[0]
                self.list.addItem(QtWidgets.QListWidgetItem(text))

    def insert(self):
        self.w = InsertWindow(self)
        self.w.setWindowTitle('Добавить')
        self.w.show()
        self.w.showMaximized()
        
    def change_data(self):
        if self.list.currentItem():
            data = self.db.get_data(self.list_items[self.list.currentItem().text()])[0]
            self.w = ChangeWindow(data, self)
            self.w.setWindowTitle('Редактировать')
            self.w.show()
            self.w.showMaximized()

    def delete_profile(self):
        if self.list.currentItem():
            data = self.db.get_data(self.list_items[self.list.currentItem().text()])[0]
            result = QtWidgets.QMessageBox.question(self,
                                                    'Удаление учетной записи',
                                                    'Удалить учетную запись {} {}.{}.?'.format(data[2], data[3][0], data[4][0]),
                                                    QtWidgets.QMessageBox.Yes | QtWidgets.QMessageBox.No,
                                                    QtWidgets.QMessageBox.No)
            if result == QtWidgets.QMessageBox.Yes:
                query = """
                DELETE FROM persons
                WHERE persons."ID" = '{}';
                """.format(data[0])
                self.db.execute(query)
                self.find()
                self.splitter.replaceWidget(1, MainTab(None, parent=self))

    def chose_person(self, idx):
        self.splitter.replaceWidget(1, MainTab(idx, parent=self))
        

    def data_change(self):
        self.w = DataWindow(self)
        self.w.setWindowTitle('Справочник')
        self.w.show()

    def fizo_open(self):
        self.w = FizoWindow(self)
        self.w.setWindowTitle('ФИЗО')
        self.w.show()
        
    def add_scan(self):
        if self.list.currentItem():
            data = self.db.get_data(self.list_items[self.list.currentItem().text()])[0][0]
            self.w = ScanWindow(data, self)
            self.w.setWindowTitle('Прикрепить документ')
            self.w.show()

    def create_documents(self):
        self.doc_w = DocWindow(self)
        self.doc_w.setWindowTitle('Документы')
        self.doc_w.show()

    def import_db(self):
        """Импорт данных из базы (.db) старой версии.

        Порядок: подтверждение -> бекап текущей базы (дата_время_backup.db) ->
        очистка текущей базы -> перенос строк из выбранного файла.
        """
        answer = QtWidgets.QMessageBox.question(
            self, 'Импорт данных',
            'Импорт данных очистит базу. Вы уверены что хотите продолжить?',
            QtWidgets.QMessageBox.Yes | QtWidgets.QMessageBox.No,
            QtWidgets.QMessageBox.No)
        if answer != QtWidgets.QMessageBox.Yes:
            return

        file_path, _ = QtWidgets.QFileDialog.getOpenFileName(
            self, 'Выберите файл базы данных старой версии', '',
            'Базы данных (*.db);;Все файлы (*)')
        if not file_path:
            return

        backup_path = None
        try:
            backup_path = self.db.backup_current_db()
            self.db.clear_db()
            imported = self.db.import_from_file(file_path)
        except Exception as e:
            QtWidgets.QMessageBox.critical(
                self, 'Ошибка импорта',
                'Не удалось выполнить импорт: {}\nРезервная копия: {}'.format(
                    e, backup_path or 'не создан'))
            return

        total = sum(imported.values())
        detail = ', '.join('{}: {}'.format(t, c) for t, c in imported.items() if c)
        QtWidgets.QMessageBox.information(
            self, 'Импорт завершён',
            'Перенесено записей: {}.\n({})\n\nРезервная копия сохранена:\n{}'
            .format(total, detail or 'данных не найдено', backup_path))

        # обновляем интерфейс после импорта
        self.splitter.replaceWidget(1, MainTab(None, parent=self))
        self.find()

    def export(self):
        f = QtWidgets.QFileDialog.getSaveFileName(self, 'Сохранить как', 'Сотрудники.xlsx')
        file_name = f[0]

        if file_name:
            wb = openpyxl.Workbook()
            wb.create_sheet('Персональные данные')
            sheet = wb['Персональные данные']
            query = f'''
            SELECT 
            persons.lnum, 
            persons.second_name, 
            persons.first_name, 
            persons.patr, 
            persons.person_buisness,
            persons.family_status,
            persons."rank",
            persons.division,
            persons.position,
            persons.order_by,
            persons.order_date,
            persons.place_of_birth, 
            persons.date_of_birth, 
            persons.gender, 
            persons.nation, 
            persons.phone_num1,
            persons.phone_num2,
            persons.address_reg,
            persons.address_living,
            persons.allowance_num,
            persons.allowance_by,
            persons.allowance_date,
            snils,
            inn,
            fizo_num,
            "ID"
            FROM persons
            WHERE persons.second_name LIKE '%' {self.search_spec_query}
            ORDER BY persons.second_name
            '''
            res = self.db.select(query)
            sheet.append(('Личный номер', 'ФИО', 'Фамилия', 'Имя', 'Отчество', 'Номер личного дела', 'Семейное положение', 'Звание', 'Подразделение',
                          'Должность', 'Приказ', 'Дата приказа',
                          'Место рождения', 'Дата рождения', 'Пол', 'Национальность',
                          'Телефон1', 'Телефон2', 'Адрес регистрации', 'Адрес проживания', 'Номер допуска', 'Кем допущен',
                          'Дата допуска', 'СНИЛС', 'ИНН', 'Номерок', 'Возрастная гр'))
            if res:
                for row in res:

                    years = datetime.datetime.now().year - row[12].year
                    if datetime.datetime.now().month < row[12].month:
                        years -= 1
                    elif datetime.datetime.now().month == row[12].month:
                        if datetime.datetime.now().day < row[12].day:
                            years -= 1

                    if years < 25:
                        group = '1'
                    elif years < 30:
                        group = '2'
                    elif years < 35:
                        group = '3'
                    elif years < 40:
                        group = '4'
                    elif years < 45:
                        group = '5'
                    elif years < 50:
                        group = '6'
                    elif years < 55:
                        group = '7'
                    else:
                        group = '8'

                    name = f'{row[1]} {row[2][0]}.{row[3][0]}.'
                    sheet.append((row[0], name,) + row[1:] + (group, ))

                wb.create_sheet('Классность')
                sheet = wb['Классность']
                for row in res:
                    idx = row[-1]
                    query = f'''
                    SELECT 
                    persons.lnum, 
                    persons.second_name, 
                    persons.first_name, 
                    persons.patr, 
                    class.level,
                    class.order_date,
                    class.order_by,
                    class.date_passing,
                    class.actual_from,
                    class.actual_to,
                    class.vus
                    FROM persons, class
                    WHERE persons.ID = class.PID AND
                    persons.ID = {idx}
                    ORDER BY class.order_date DESC
                    '''
                    res = self.db.select(query)
                    if res:
                        row = res[0]
                        sheet.append(row)

            wb.create_sheet('Документы')
            sheet = wb['Документы']
            query = '''
                    SELECT 
                    persons.lnum, 
                    persons.second_name, 
                    persons.first_name, 
                    persons.patr, 
                    documents."type",
                    documents.serial,
                    documents.num,
                    documents.given_by,
                    documents.code,
                    documents.doc_date
                    FROM persons, documents
                    WHERE persons.ID = documents.PID
                    ORDER BY persons.second_name
                    '''
            res = self.db.select(query)
            if res:
                for row in res:
                    sheet.append(row)

            wb.create_sheet('Члены семьи')
            sheet = wb['Члены семьи']
            query = '''
                    SELECT 
                    persons.lnum, 
                    persons.second_name, 
                    persons.first_name, 
                    persons.patr, 
                    family."type",
                    family.second_name,
                    family.first_name,
                    family.patr,
                    family.birthday
                    FROM persons, family
                    WHERE persons.ID = family.PID
                    ORDER BY persons.second_name
                    '''
            res = self.db.select(query)
            if res:
                for row in res:
                    sheet.append(row)

            wb.create_sheet('Ипотека')
            sheet = wb['Ипотека']
            query = '''
                SELECT 
                persons.lnum, 
                persons.second_name, 
                persons.first_name, 
                persons.patr, 
                ipoteka."user",
                ipoteka.finished,
                ipoteka.number,
                ipoteka.date_joining
                FROM persons, ipoteka
                WHERE persons.ID = ipoteka.PID
                ORDER BY persons.second_name
                '''
            res = self.db.select(query)
            if res:
                for row in res:
                    sheet.append(row)

            wb.create_sheet('Кредиты')
            sheet = wb['Кредиты']
            query = '''
                SELECT 
                persons.lnum, 
                persons.second_name, 
                persons.first_name, 
                persons.patr, 
                dolg.credit_org,
                dolg.credit_type,
                dolg.summ
                FROM persons, dolg
                WHERE persons.ID = dolg.PID
                ORDER BY persons.second_name
                '''
            res = self.db.select(query)
            if res:
                for row in res:
                    sheet.append(row)

            wb.create_sheet('Авто')
            sheet = wb['Авто']
            query = '''
                    SELECT 
                    persons.lnum, 
                    persons.second_name, 
                    persons.first_name, 
                    persons.patr, 
                    auto.model,
                    auto.gos_num,
                    auto.agency_name,
                    auto.agency_num,
                    auto.agency_from,
                    auto.agency_to
                    FROM persons, auto
                    WHERE persons.ID = auto.PID
                    ORDER BY persons.second_name
    
                    '''
            res = self.db.select(query)
            if res:
                for row in res:
                    sheet.append(row)

            wb.save(file_name)



if __name__ == '__main__':
    import sys

    app = QtWidgets.QApplication(sys.argv)

    window = MainWindow('employees_center')
    window.resize(1200, 600)
    # исправлен путь к иконке приложения (было './data.logo.png')
    ico = QtGui.QIcon('./data/logo.png')
    window.setWindowTitle('СПО учёта личного состава')
    window.setWindowIcon(ico)
    window.showMaximized()
    window.show()
    
    app.setWindowIcon(ico)

    sys.exit(app.exec_())
    
    
